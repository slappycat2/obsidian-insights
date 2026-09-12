"""End-to-end: the real SysConfig, the real CLI, no Obsidian and no window.

Everything else in the suite goes through StubSysConfig, which is why the
platform bugs in SysConfig survived a green CI on all three runners. These
tests run ``ovi --headless --do-not-open`` in a child process whose home
directory is an empty temp folder -- so there is no obsidian.json to find --
and whose spreadsheet application is blank. That is exactly the state of a
fresh install on a Mac or a Linux box, and it used to fail before the setup
screen could open.

A headless run cannot show the setup screen, so on a machine with no
CONFIG.yaml it builds one from the defaults when -- and only when -- a
VAULT_PATH names what to scan. That is the path the Obsidian plugin takes on
its first run. With no vault there is nothing to build a config around, and
the run still stops with advice.
"""

import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
import yaml

from ovi import CTOT_SLOTS, __version__

REPO_ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture
def clean_env(tmp_path):
    """A child-process environment with nothing of this machine's in it."""
    home = tmp_path / "home"
    home.mkdir()
    data = tmp_path / "data"
    env = dict(os.environ)
    env.update({
        "OVI_DATA_DIR": str(data),
        "HOME": str(home),
        "USERPROFILE": str(home),
        "APPDATA": str(home / "AppData" / "Roaming"),
        "XDG_CONFIG_HOME": str(home / ".config"),
        "PYTHONIOENCODING": "utf-8",
    })
    return env, data


def run(args, env):
    return subprocess.run([sys.executable, *args], env=env, cwd=REPO_ROOT,
                          capture_output=True, text=True, encoding="utf-8", timeout=120)


def test_headless_run_with_no_obsidian_and_no_spreadsheet_app(clean_env, make_vault):
    """First run on a fresh machine: no config, no obsidian.json, a vault path.

    The config is created from the defaults and the workbook is written, all
    without a window. This used to need a SysConfig subclass whose
    run_setup_ui() saved instead of opening the screen.
    """
    env, data = clean_env
    vault = make_vault({"Note.md": "---\ntitle: Hello\ntags: [a]\n---\nBody #tag\n"},
                       name="Fresh Vault")

    result = run(["-m", "ovi.ovi", "--headless", "--do-not-open", "-d", "WARNING", str(vault)], env)

    assert result.returncode == 0, result.stderr
    assert "Traceback" not in result.stderr
    assert (data / "CONFIG.yaml").is_file(), "the config was not bootstrapped"
    workbooks = list((data / "data" / "workbooks").glob("ovi_Fresh_Vault_*.xlsx"))
    assert len(workbooks) == 1, result.stdout + result.stderr
    assert result.stdout.strip() == f"Workbook written to {workbooks[0]}"

    saved = yaml.safe_load((data / "CONFIG.yaml").read_text(encoding="utf-8"))
    assert saved["dir_vault"] == str(vault)
    # The platform default: a spreadsheet program found in its usual place, or
    # blank for the system handler. Either way it has to validate, because the
    # next run's chk_fields_on_load() gates on it.
    from ovi.ovi_setup import SysConfig
    assert SysConfig.validate_sys_pn_wb_exec(saved["sys_pn_wb_exec"])[0]


def test_second_headless_run_reuses_the_bootstrapped_config(clean_env, make_vault):
    """The bootstrap is a one-off; the next run finds a valid config and
    numbers its workbook after the first."""
    env, data = clean_env
    vault = make_vault({"Note.md": "text\n"}, name="Fresh Vault")
    args = ["-m", "ovi.ovi", "--headless", "--do-not-open", "-d", "WARNING", str(vault)]

    first = run(args, env)
    stamp = (data / "CONFIG.yaml").stat().st_mtime_ns
    second = run(args, env)

    assert first.returncode == 0 and second.returncode == 0, first.stderr + second.stderr
    assert (data / "CONFIG.yaml").stat().st_mtime_ns == stamp, "a plain run rewrote the config"
    names = sorted(p.name for p in (data / "data" / "workbooks").glob("*.xlsx"))
    assert names == ["ovi_Fresh_Vault_0000.xlsx", "ovi_Fresh_Vault_0001.xlsx"]


def test_headless_run_without_setup_or_a_vault_says_what_to_do(clean_env):
    """No CONFIG.yaml, no VAULT_PATH and --headless: a one-line error, exit 1,
    no traceback -- and the message names both ways out."""
    env, _ = clean_env

    result = run(["-m", "ovi.ovi", "--headless", "--do-not-open"], env)

    assert result.returncode == 1
    assert "Traceback" not in result.stderr
    assert "--headless" in result.stderr
    assert "VAULT_PATH" in result.stderr


def test_setup_flag_is_never_satisfied_by_a_bootstrap(clean_env, make_vault):
    """--setup asks for the screen; headless cannot show it, and must not
    quietly write a config instead."""
    env, data = clean_env
    vault = make_vault({"Note.md": "text\n"})

    result = run(["-m", "ovi.ovi", "--headless", "--setup", "--do-not-open", str(vault)], env)

    assert result.returncode == 1
    assert "Traceback" not in result.stderr
    assert not (data / "CONFIG.yaml").exists()


# ---------------------------------------------------------------------------
# --json, end to end: what the Obsidian plugin reads
# ---------------------------------------------------------------------------

def json_events(result):
    lines = result.stdout.splitlines()
    assert lines, "nothing on stdout: " + result.stderr
    return [json.loads(line) for line in lines]


def test_json_mode_emits_progress_then_done(clean_env, make_vault):
    env, data = clean_env
    vault = make_vault({"Note.md": "---\ntitle: Hello\n---\nBody #tag\n"}, name="Fresh Vault")

    result = run(["-m", "ovi.ovi", "--json", "-x", "-d", "WARNING", str(vault)], env)

    assert result.returncode == 0, result.stderr
    got = json_events(result)
    assert [e["event"] for e in got] == ["progress"] * 5 + ["done"]
    assert [e["percent"] for e in got[:-1]] == [10, 20, 50, 70, 100]
    done = got[-1]
    assert done["ok"] is True
    assert Path(done["workbook"]).is_file()
    assert Path(done["batch"]).is_file()
    assert done["vault"] == str(vault)
    assert done["version"] == __version__
    assert len(done["ctot"]) == CTOT_SLOTS
    assert done["ctot"][0] == 1, "one markdown file seen"
    assert done["opened"] is None, "-x leaves the workbook unopened"
    assert "Workbook written to" not in result.stdout
    assert (data / "CONFIG.yaml").is_file(), "--json implies --headless, which bootstraps"


def test_json_mode_reports_a_missing_config_as_one_line(clean_env):
    env, _ = clean_env

    result = run(["-m", "ovi.ovi", "--json", "-x"], env)

    assert result.returncode == 1
    got = json_events(result)
    assert len(got) == 1
    assert got[0]["event"] == "error"
    assert got[0]["kind"] == "ConfigIncomplete"
    assert "VAULT_PATH" in got[0]["message"]
    assert "Traceback" not in result.stderr


def test_override_flags_reach_the_workbook(clean_env, make_vault):
    """The settings the plugin passes have to change the output, not just the
    config: a skipped folder is not scanned, and the link limits cap the
    FileNN columns on the Values and Tags tabs."""
    import openpyxl

    env, data = clean_env
    note = "---\nstatus: x\n---\nBody #t\n"
    vault = make_vault({"Notes/A.md": note, "Notes/B.md": note, "Notes/C.md": note,
                        "Archive/Old.md": note}, name="Vault")

    result = run(["-m", "ovi.ovi", "--json", "-x", "-d", "WARNING",
                  "--skip-folders", "Archive", "--max-value-links", "1", "--max-tag-links", "1",
                  str(vault)], env)

    assert result.returncode == 0, result.stderr
    done = json_events(result)[-1]
    assert done["ctot"][2] == 1, "one file skipped by folder"
    assert done["ctot"][3] == 3, "three files analysed"

    wb = openpyxl.load_workbook(done["workbook"])
    for tab_name in ("Values", "Tags"):
        headers = [cell.value for row in wb[tab_name].iter_rows(min_row=1, max_row=12)
                   for cell in row if isinstance(cell.value, str) and cell.value.startswith("File")]
        assert "File01" in headers, f"{tab_name}: no link columns at all"
        assert "File02" not in headers, f"{tab_name}: the link limit was ignored"
    text = " ".join(str(c.value) for row in wb["Values"].iter_rows() for c in row if c.value)
    assert "Old" not in text, "the skipped folder's note reached the Values tab"

    saved = yaml.safe_load((data / "CONFIG.yaml").read_text(encoding="utf-8"))
    assert saved["cur_vlts"][done["vault_name"]]["skip_rel_str"] == "Archive"


def test_an_invalid_spreadsheet_app_is_a_config_error(clean_env, make_vault, tmp_path):
    env, data = clean_env
    vault = make_vault({"Note.md": "text\n"})

    result = run(["-m", "ovi.ovi", "--json", "-x", "--spreadsheet-app", str(tmp_path / "nope.exe"),
                  str(vault)], env)

    assert result.returncode == 1
    got = json_events(result)
    assert got[-1]["kind"] == "ConfigIncomplete"
    assert "nope.exe" in got[-1]["message"]
    assert not (data / "CONFIG.yaml").exists()


def test_a_blank_spreadsheet_app_means_the_system_default(clean_env, make_vault):
    env, data = clean_env
    vault = make_vault({"Note.md": "text\n"})

    result = run(["-m", "ovi.ovi", "--json", "-x", "--spreadsheet-app", "", str(vault)], env)

    assert result.returncode == 0, result.stderr
    assert json_events(result)[-1]["event"] == "done"
    saved = yaml.safe_load((data / "CONFIG.yaml").read_text(encoding="utf-8"))
    assert saved["sys_pn_wb_exec"] == ""
